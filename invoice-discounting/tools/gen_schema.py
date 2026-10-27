import openpyxl, json, sys, re
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)

def txt(c):
    return None if c.value is None else str(c.value).replace('\r', '').strip()

def mand_kind(m):
    m = (m or '').strip()
    if m.startswith('Yes*'): return 'conditional'
    if m.startswith('Yes'): return 'always'
    if m.startswith('No'): return 'optional'
    return 'unspecified'

def parse(ws):
    res = {'request': [], 'response': []}
    cur = None
    for row in ws.iter_rows(min_row=1, max_col=9):
        a = txt(row[0])
        if a and re.search(r'(Request|Response)$', a) and not row[1].value:
            cur = 'request' if a.endswith('Request') else 'response'
            continue
        if cur is None or a == 'Sr. No.' or not (row[3].value or row[2].value):
            continue
        g = txt(row[6]) or ''
        res[cur].append({
            'name': txt(row[1]), 'path': (txt(row[2]) or '') if row[3].value else '', 'key': txt(row[3]) or txt(row[2]),
            'desc': txt(row[4]) or '', 'type': txt(row[5]) or '',
            'mand': g.split('\n')[0].strip(), 'mandRule': g, 'mandKind': mand_kind(g),
            'sample': txt(row[7]) or '', 'validation': txt(row[8]) or '',
        })
    return res

schema = {
    'leadApi': parse(wb['Lead API']),
    'getStatus': parse(wb['Get Status API']),
    'partnerships': parse(wb['Get Partnerships Details API']),
    'channelPartners': parse(wb['Get Channel Partners for a Prog']),
    'additionalLenders': parse(wb['Get Additional Lenders on PSB A']),
}

# masters
ws = wb['Master Reference']
masters = {'leadTypes': [], 'onboardingModes': [], 'products': [], 'conduct': [], 'constitution': []}
sect = 'leadTypes'
for row in ws.iter_rows(min_row=1, max_col=4):
    v = [txt(c) for c in row]
    if not any(v): continue
    if v[0] == 'Lead Type Code': sect = 'leadTypes'; continue
    if v[0] == 'Lead Onboarding Mode': sect = 'onboardingModes'; continue
    if v[0] == 'Product Master': sect = 'products'; continue
    if v[0] and v[0].startswith('Conduct Master'):
        sect = 'conduct' if sect != 'conduct' else 'constitution'; continue
    if v[0] in ('Sr. No.', 'Product Code'): continue
    if sect == 'leadTypes' and v[0]: masters[sect].append({'code': v[0], 'name': v[1], 'remarks': v[2] or ''})
    elif sect == 'onboardingModes': masters[sect].append({'name': v[1], 'remarks': v[2] or ''})
    elif sect == 'products': masters[sect].append({'code': v[0], 'name': v[1]})
    elif sect == 'conduct': masters[sect].append({'code': v[1], 'desc': v[2] or ''})
    elif sect == 'constitution': masters[sect].append({'code': v[1], 'desc': v[2] or ''})

ws = wb['Status Master']
status = {'fintech': [], 'lender': []}
sect = None
for row in ws.iter_rows(min_row=1, max_col=5):
    v = [txt(c) for c in row]
    if not any(v): continue
    if v[0] == 'Fintech Status Master': sect = 'fintech'; continue
    if v[0] == 'Lender Status Master': sect = 'lender'; continue
    if v[0] == 'Sr. No.' or sect is None: continue
    status[sect].append({'stage': v[1], 'status': v[2], 'desc': v[3] or ''})

doc = ws = wb['Document Description']
apis = []
for row in doc.iter_rows(min_row=9, max_col=6):
    v = [txt(c) for c in row]
    if v[1]: apis.append({'name': v[1], 'desc': v[2], 'type': v[3], 'exposedBy': v[4]})

js = ('/* GENERATED from "PSBPlatform Xchange_Sourcing_Partner_Integration_API_Documentation_Schema_v1.1.xlsx" (API version 1.1, 2024-05-20).\n'
      '   Do not hand-edit: re-run the extractor. Field rows, mandatory flags and validation text are verbatim from the sheet. */\n'
      'const PSB_SCHEMA = ' + json.dumps(schema, ensure_ascii=False, indent=1) + ';\n'
      'const PSB_MASTERS = ' + json.dumps(masters, ensure_ascii=False, indent=1) + ';\n'
      'const PSB_STATUS_MASTER = ' + json.dumps(status, ensure_ascii=False, indent=1) + ';\n'
      'const PSB_API_INDEX = ' + json.dumps(apis, ensure_ascii=False, indent=1) + ';\n')
open(out, 'w', encoding='utf-8').write(js)
for k, v in schema.items(): print(k, len(v['request']), len(v['response']))
print({k: len(v) for k, v in masters.items()}, {k: len(v) for k, v in status.items()}, len(apis))

